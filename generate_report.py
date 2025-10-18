import os
from datetime import datetime
from typing import Dict

import pandas as pd

from src.analysis import load_netflix_dataset, clean_and_engineer, compute_views, compute_insights, export_key_charts, ensure_reports_dir

# Optional dependencies are imported lazily below


def ensure_optional_packages():
    missing = []
    try:
        from pptx import Presentation  # noqa: F401
    except Exception:
        missing.append('python-pptx')
    try:
        from reportlab.pdfgen import canvas  # noqa: F401
        from reportlab.lib.pagesizes import letter  # noqa: F401
    except Exception:
        missing.append('reportlab')
    if missing:
        print('Note: Missing optional packages ->', ', '.join(missing))
        print('Install them to enable PPTX/PDF export:')
        print('  pip install python-pptx reportlab')


def build_powerpoint(chart_paths: Dict[str, str], insights, out_path: str):
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = 'Netflix Content Trends — Summary'
    slide.placeholders[1].text = f'Generated on {datetime.now():%Y-%m-%d %H:%M}'

    # Insights slide
    bullet_layout = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(bullet_layout)
    slide2.shapes.title.text = 'Key Insights'
    tf = slide2.placeholders[1].text_frame
    tf.clear()
    for line in insights[:6]:
        p = tf.add_paragraph()
        p.text = line
        p.level = 0

    # Charts slides
    for name, path in chart_paths.items():
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = name.replace('_', ' ').title()
        left = Inches(1)
        top = Inches(1.5)
        height = Inches(5)
        slide.shapes.add_picture(path, left, top, height=height)

    prs.save(out_path)


def build_pdf(chart_paths: Dict[str, str], insights, out_path: str):
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.utils import ImageReader

    c = canvas.Canvas(out_path, pagesize=letter)
    width, height = letter

    # Title page
    c.setFont('Helvetica-Bold', 20)
    c.drawString(72, height - 72, 'Netflix Content Trends — Summary')
    c.setFont('Helvetica', 12)
    c.drawString(72, height - 96, f'Generated on {datetime.now():%Y-%m-%d %H:%M}')
    c.showPage()

    # Insights page
    c.setFont('Helvetica-Bold', 16)
    c.drawString(72, height - 72, 'Key Insights')
    c.setFont('Helvetica', 12)
    y = height - 100
    for line in insights[:12]:
        c.drawString(72, y, f'- {line}')
        y -= 18
        if y < 72:
            c.showPage()
            y = height - 72
    c.showPage()

    # Charts pages
    for name, path in chart_paths.items():
        c.setFont('Helvetica-Bold', 16)
        c.drawString(72, height - 72, name.replace('_', ' ').title())
        try:
            img = ImageReader(path)
            # Fit image within page margins
            max_w, max_h = width - 144, height - 144
            c.drawImage(img, 72, 72, width=max_w, height=max_h, preserveAspectRatio=True, anchor='c')
        except Exception as e:
            c.setFont('Helvetica', 12)
            c.drawString(72, height - 120, f'(Image not available: {e})')
        c.showPage()

    c.save()


def main():
    ensure_optional_packages()

    reports_dir = ensure_reports_dir('reports')

    # Strict local-only dataset usage; auto-detects 'Netflix Dataset.csv' or other local 'netflix*.csv'
    df_raw = load_netflix_dataset(None, strict=True)
    df = clean_and_engineer(df_raw)
    analysis = compute_views(df)

    # Export charts
    chart_paths = export_key_charts(analysis, reports_dir=reports_dir)

    # Build insights
    insights = compute_insights(analysis)

    # Create PPTX and PDF
    ts = datetime.now().strftime('%Y%m%d_%H%M')
    pptx_out = os.path.join(reports_dir, f'netflix_summary_{ts}.pptx')
    pdf_out = os.path.join(reports_dir, f'netflix_summary_{ts}.pdf')

    try:
        build_powerpoint(chart_paths, insights, pptx_out)
        print('PowerPoint saved:', pptx_out)
    except Exception as e:
        print('PPTX generation skipped or failed:', e)

    try:
        build_pdf(chart_paths, insights, pdf_out)
        print('PDF saved:', pdf_out)
    except Exception as e:
        print('PDF generation skipped or failed:', e)


if __name__ == '__main__':
    main()
